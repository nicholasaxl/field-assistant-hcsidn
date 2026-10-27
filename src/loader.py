"""Document loader: turns files (bulk folder or single upload) into normalized Documents.

Each loader maps a file to a list of LangChain Documents. `load_file` then stamps the
common metadata (source, doc_type, file_hash, device_model). To support a new format,
write one `_load_xxx(path)` function and add it to LOADERS.
"""
from __future__ import annotations

import csv
import hashlib
import re
from pathlib import Path
from typing import Callable, Iterator

from langchain_core.documents import Document

import config


class UnsupportedFileType(ValueError):
    pass


class InvalidUpload(ValueError):
    pass


class EmptyDocument(ValueError):
    """The file parsed but contained no extractable text (e.g. a scanned PDF)."""


_FRONT_MATTER = re.compile(r"\A---\s*\n(.*?)\n---\s*\n", re.DOTALL)
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$", re.MULTILINE)
CSV_ROWS_PER_DOC = 20


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def _split_front_matter(text: str) -> tuple[dict, str]:
    match = _FRONT_MATTER.match(text)
    if not match:
        return {}, text
    meta = {}
    for line in match.group(1).splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            meta[key.strip().lower()] = value.strip()
    return meta, text[match.end():]


def _load_pdf(path: Path) -> list[Document]:
    from pypdf import PdfReader

    docs = []
    for number, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            docs.append(Document(page_content=text, metadata={"page": number}))
    return docs


def _load_docx(path: Path) -> list[Document]:
    import docx

    document = docx.Document(str(path))
    sections: list[tuple[str, list[str]]] = [("", [])]
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name.lower() if paragraph.style is not None else ""
        if style.startswith("heading") or style == "title":
            sections.append((text, []))
        else:
            sections[-1][1].append(text)
    for number, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        sections.append((f"Table {number}", rows))
    return [
        Document(page_content="\n".join(lines), metadata={"section": title} if title else {})
        for title, lines in sections
        if lines
    ]


def _load_markdown(path: Path) -> list[Document]:
    """Markdown / plain text. Optional `---` front matter (e.g. device_model: HX-POS-200)."""
    raw = path.read_text(encoding="utf-8", errors="replace")
    front, body = _split_front_matter(raw)
    base = {"device_model": front["device_model"]} if front.get("device_model") else {}

    matches = list(_HEADING.finditer(body))
    if not matches:
        text = body.strip()
        return [Document(page_content=text, metadata=dict(base))] if text else []

    docs = []
    preamble = body[: matches[0].start()].strip()
    if preamble:
        docs.append(Document(page_content=preamble, metadata=dict(base)))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(body)
        text = body[match.end():end].strip()
        if text:
            docs.append(Document(page_content=text, metadata={**base, "section": match.group(2).strip()}))
    return docs


def _load_csv(path: Path) -> list[Document]:
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
        rows = list(csv.DictReader(fh))
    docs = []
    for start in range(0, len(rows), CSV_ROWS_PER_DOC):
        batch = rows[start:start + CSV_ROWS_PER_DOC]
        lines = ["; ".join(f"{k}: {v}" for k, v in row.items() if v) for row in batch]
        docs.append(
            Document(
                page_content="\n".join(lines),
                metadata={"section": f"rows {start + 1}-{start + len(batch)}"},
            )
        )
    return docs


LOADERS: dict[str, Callable[[Path], list[Document]]] = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".md": _load_markdown,
    ".markdown": _load_markdown,
    ".txt": _load_markdown,
    ".csv": _load_csv,
}
SUPPORTED_EXTENSIONS = sorted(LOADERS)


def load_file(path: str | Path, device_model: str = config.ALL_DEVICES) -> list[Document]:
    """Load one file into Documents with normalized metadata."""
    path = Path(path)
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise UnsupportedFileType(
            f"'{path.suffix or path.name}' is not supported. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    docs = loader(path)
    if not docs:
        raise EmptyDocument(f"No extractable text found in '{path.name}' (scanned PDFs need OCR).")

    digest = file_hash(path)
    for doc in docs:
        doc.metadata.setdefault("device_model", device_model or config.ALL_DEVICES)
        doc.metadata.update(
            source=path.name,
            doc_type=path.suffix.lower().lstrip("."),
            file_hash=digest,
        )
    return docs


def iter_supported_files(directory: str | Path) -> Iterator[Path]:
    for path in sorted(Path(directory).rglob("*")):
        if path.is_file() and path.suffix.lower() in LOADERS:
            yield path


def validate_upload(filename: str, size_bytes: int) -> str:
    """Check type and size of an uploaded file; return a filesystem-safe filename."""
    name = Path(filename).name
    if Path(name).suffix.lower() not in LOADERS:
        raise InvalidUpload(f"'{name}': unsupported type. Allowed: {', '.join(SUPPORTED_EXTENSIONS)}")
    if size_bytes > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise InvalidUpload(f"'{name}' is larger than {config.MAX_UPLOAD_MB} MB")
    if size_bytes == 0:
        raise InvalidUpload(f"'{name}' is empty")
    return re.sub(r"[^\w.\- ]", "_", name)
