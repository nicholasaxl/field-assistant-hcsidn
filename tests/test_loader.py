import docx
import pytest

import config
from src.loader import (
    EmptyDocument,
    InvalidUpload,
    UnsupportedFileType,
    iter_supported_files,
    load_file,
    validate_upload,
)


def test_markdown_front_matter_and_sections(tmp_path):
    path = tmp_path / "runbook.md"
    path.write_text("---\ndevice_model: HX-POS-200\n---\n# Jam\nClear the paper.\n## Reload\nLoad a roll.\n", encoding="utf-8")
    docs = load_file(path)
    assert [d.metadata["section"] for d in docs] == ["Jam", "Reload"]
    assert all(d.metadata["device_model"] == "HX-POS-200" for d in docs)
    assert all(d.metadata["source"] == "runbook.md" and d.metadata["doc_type"] == "md" for d in docs)
    assert len({d.metadata["file_hash"] for d in docs}) == 1


def test_device_model_defaults_to_argument(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Check the cable.", encoding="utf-8")
    assert load_file(path, "HX-ATM-500")[0].metadata["device_model"] == "HX-ATM-500"
    assert load_file(path)[0].metadata["device_model"] == config.ALL_DEVICES


def test_csv_rows_are_grouped(tmp_path):
    path = tmp_path / "codes.csv"
    rows = "\n".join(f"E-{1000 + i},issue {i}" for i in range(25))
    path.write_text("code,meaning\n" + rows, encoding="utf-8")
    docs = load_file(path)
    assert len(docs) == 2
    assert "code: E-1000; meaning: issue 0" in docs[0].page_content


def test_docx_headings_and_tables(tmp_path):
    path = tmp_path / "manual.docx"
    document = docx.Document()
    document.add_heading("Printer", level=1)
    document.add_paragraph("Replace the roll.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Code"
    table.rows[0].cells[1].text = "P-1001"
    document.save(path)
    docs = load_file(path)
    sections = {d.metadata.get("section"): d.page_content for d in docs}
    assert "Replace the roll." in sections["Printer"]
    assert "P-1001" in sections["Table 1"]


def _minimal_pdf(text: bytes) -> bytes:
    stream = b"BT /F1 12 Tf 72 720 Td (" + text + b") Tj ET"
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer<</Root 1 0 R/Size %d>>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref_at)
    return out


def test_pdf_pages(tmp_path):
    path = tmp_path / "manual.pdf"
    path.write_bytes(_minimal_pdf(b"Replace the printer paper roll"))
    docs = load_file(path)
    assert docs[0].metadata["page"] == 1
    assert "printer paper roll" in docs[0].page_content


def test_unsupported_and_empty(tmp_path):
    exe = tmp_path / "tool.exe"
    exe.write_bytes(b"MZ")
    with pytest.raises(UnsupportedFileType):
        load_file(exe)
    empty = tmp_path / "blank.md"
    empty.write_text("   \n", encoding="utf-8")
    with pytest.raises(EmptyDocument):
        load_file(empty)


def test_iter_supported_files_skips_unknown(tmp_path):
    (tmp_path / "a.md").write_text("x", encoding="utf-8")
    (tmp_path / "b.exe").write_bytes(b"x")
    assert [p.name for p in iter_supported_files(tmp_path)] == ["a.md"]


def test_validate_upload():
    assert validate_upload("../../evil name!.md", 10) == "evil name_.md"
    with pytest.raises(InvalidUpload):
        validate_upload("x.exe", 10)
    with pytest.raises(InvalidUpload):
        validate_upload("x.md", config.MAX_UPLOAD_MB * 1024 * 1024 + 1)
    with pytest.raises(InvalidUpload):
        validate_upload("x.md", 0)
